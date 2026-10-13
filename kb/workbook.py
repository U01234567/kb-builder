"""kb.xlsx: read and write questions, answers, evidence and reviews, and keep it readable.

Topic sheets are made by people. kb-builder makes and owns five fixed sheets. Columns are found by
their header, never by position. Derived columns are rewritten on every save, so they cannot drift.
"""

import hashlib
import os
import re
import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from kb.index import by_content

TOPICS, OVERVIEW, EVIDENCE, SOURCES, RUNS = "Topics", "Overview", "Evidence", "Sources", "Runs"
FIXED = {
    OVERVIEW: ("Topic", "Reads", "Questions", "Answered", "Accept", "Revisions", "Reject", "Not reviewed"),
    TOPICS: ("Topic", "Reads"),
    EVIDENCE: ("Question", "Source", "Title", "File", "Location", "Quote", "Verified", "Removed", "Reason"),
    SOURCES: ("ID", "Files", "Title", "URL", "Saved on", "Fingerprint", "Status"),
    RUNS: ("Started", "Finished", "Machine", "Mode", "Topics", "Model", "Embeddings", "Prompts", "Questions"),
}
STATUSES = ("Accept", "Revisions", "Reject")
DERIVED_NOTE = "Derived by kb-builder from other sheets. Edits here are overwritten on the next save."


class ChangedOnDisk(Exception):
    """kb.xlsx changed since it was read: someone else saved it. Nothing was written."""


# ---------------------------------------------------------------- opening and saving
def stamp(path: Path) -> str:
    """A hash of the file as it is on disk now, to notice later whether someone else saved it."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


def open_kb(path: Path) -> tuple[Workbook, str]:
    """The workbook, and the stamp to pass to save()."""
    return load_workbook(path), stamp(path)


def save(wb: Workbook, path: Path, expected: str, project) -> str:
    """Update derived columns and layout, then write through a temporary file. Returns the new stamp.

    Refuses (ChangedOnDisk) when the file on disk is no longer the one that was opened.
    """
    if stamp(path) != expected:
        raise ChangedOnDisk(f"{path.name} was changed by someone else since it was opened. Open it again.")
    derive_sources(wb, project)
    derive_overview(wb, project)
    format_all(wb, project)
    temporary = path.with_name(f".{path.name}.tmp")
    wb.save(temporary)
    os.replace(temporary, path)
    return stamp(path)


# ---------------------------------------------------------------- reading
def headers(ws) -> dict[str, int]:
    """Header text -> column number, from the first row."""
    return {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value not in (None, "")}


def records(ws) -> list[dict]:
    """Each non-empty row as {header: value}, plus its row number under '_row'."""
    columns = headers(ws)
    rows = []
    for number in range(2, ws.max_row + 1):
        row = {name: ws.cell(number, column).value for name, column in columns.items()}
        if any(value not in (None, "") for value in row.values()):
            rows.append({**row, "_row": number})
    return rows


def put(ws, row: int, header: str, value) -> None:
    """Write `value` in the column with this header."""
    ws.cell(row, headers(ws)[header]).value = value


def add_row(ws, values: dict) -> int:
    """Write a new row of {header: value} below the last one and return its number.

    An empty last row (left by the table layout of an otherwise empty sheet) is used, not skipped.
    """
    row = ws.max_row + 1 if any(c.value for c in ws[ws.max_row]) else ws.max_row
    for header, value in values.items():
        put(ws, row, header, value)
    return row


def topic_sheets(wb: Workbook) -> list[str]:
    """The sheets people made: every sheet that is not one of kb-builder's five."""
    return [name for name in wb.sheetnames if name not in FIXED]


def topics(wb: Workbook) -> dict[str, str]:
    """Topic -> Reads, as written in the Topics sheet (empty when there is no Topics sheet yet)."""
    if TOPICS not in wb.sheetnames:
        return {}
    return {str(r["Topic"]).strip(): str(r.get("Reads") or "").strip() for r in records(wb[TOPICS]) if r.get("Topic")}


def questions(wb: Workbook, project) -> list[dict]:
    """Every question row of every topic sheet, keyed by column key, with '_sheet' and '_row'."""
    columns = project.topic_columns()
    rows = []
    for sheet in topic_sheets(wb):
        for r in records(wb[sheet]):
            if r.get(columns["question"]):
                rows.append({key: r.get(header) for key, header in columns.items()} | {"_sheet": sheet, "_row": r["_row"]})
    return rows


# ---------------------------------------------------------------- small pure helpers
def fingerprint(text: str) -> str:
    """A short hash of a text, to notice later that it changed."""
    return hashlib.sha256(text.encode()).hexdigest()[:16]


def slug(name: str) -> str:
    """'Opening Hours' -> 'opening-hours': the folder name kb-builder guesses for a topic."""
    plain = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", plain.lower()).strip("-")


def read_folders(reads: str, folders: set[str]) -> set[str]:
    """The folders a Reads cell names: '*' means all `folders`, otherwise names separated by commas."""
    if reads.strip() == "*":
        return set(folders)
    return {name.strip() for name in reads.split(",") if name.strip()}


def next_id(prefix: str, existing) -> str:
    """The number after the highest one in use: next_id('Q', ['Q001', 'Q009']) == 'Q010'."""
    numbers = [int(v[1:]) for v in existing if isinstance(v, str) and re.fullmatch(rf"{prefix}\d+", v)]
    return f"{prefix}{max(numbers, default=0) + 1:03d}"


# ---------------------------------------------------------------- writing
def prepare(wb: Workbook, project, folders: set[str]) -> None:
    """Make the workbook complete before a run.

    The five fixed sheets exist, every topic sheet has every column, every question has a number,
    and every topic has a row in Topics (with the guessed folder when that folder exists).
    """
    for name, columns in FIXED.items():
        ensure_columns(wb[name] if name in wb.sheetnames else wb.create_sheet(name), columns)
    labels = project.topic_columns()
    for sheet in topic_sheets(wb):
        ensure_columns(wb[sheet], labels.values())
    rows = questions(wb, project)
    used = [q["id"] for q in rows]
    for q in rows:
        if not q["id"]:
            q["id"] = next_id("Q", used)
            used.append(q["id"])
            put(wb[q["_sheet"]], q["_row"], labels["id"], q["id"])
    known = topics(wb)
    for sheet in topic_sheets(wb):
        if sheet not in known:
            guess = slug(sheet)
            add_row(wb[TOPICS], {"Topic": sheet, "Reads": guess if guess in folders else ""})


def ensure_columns(ws, wanted) -> None:
    """Add each wanted header that is missing, after the last column."""
    present = headers(ws)
    for header in wanted:
        if header not in present:
            column = max(present.values(), default=0) + 1
            ws.cell(1, column).value = header
            present[header] = column


def apply_fix(wb: Workbook, fix: tuple) -> None:
    """Apply a fix the preflight proposed: ('rename_topic', old, new) or ('set_reads', topic, reads)."""
    kind, topic, value = fix
    for r in records(wb[TOPICS]):
        if str(r["Topic"]).strip() == topic:
            put(wb[TOPICS], r["_row"], "Topic" if kind == "rename_topic" else "Reads", value)


def write_answer(wb: Workbook, project, q: dict, result, source_ids: dict, run: str, scope: str) -> None:
    """Store one generated answer with its evidence. The question's old evidence and review go.

    `source_ids` maps file fingerprints to source numbers; `scope` is a fingerprint of the files the
    topic read, so the preflight can tell later that they changed.
    """
    labels = project.topic_columns()
    values = {"answer": result.answer, "answer_simple": result.answer_simple, "run": run,
              "asked": fingerprint(q["question"]), "scope": scope,
              "status": None, "comment": None, "final_answer": None, "reviewer": None, "reviewed_at": None}
    for n in range(1, project.question_variants + 1):
        values[f"question_variant_{n}"] = result.question_variants[n - 1] if n <= len(result.question_variants) else None
    for key, value in values.items():
        put(wb[q["_sheet"]], q["_row"], labels[key], value)
    evidence = wb[EVIDENCE]
    for r in reversed(records(evidence)):
        if r["Question"] == q["id"]:
            evidence.delete_rows(r["_row"])
    for candidate, quote in result.evidence:
        add_row(evidence, {"Question": q["id"], "Source": source_ids[candidate.fingerprint],
                           "Location": candidate.location, "Quote": quote})


def sync_sources(wb: Workbook, files: dict[str, str], documents: dict, saved_on: dict) -> dict[str, str]:
    """Bring the Sources sheet in line with the files on disk, and return fingerprint -> source number.

    A source keeps its number, title and URL when its file moves (found by content), and when its
    file gets new content at the same path (found by path). New content elsewhere is a new source.
    A source whose content is gone stays, marked 'missing'. Titles and URLs people corrected are kept.
    """
    ws = wb[SOURCES]
    rows = records(ws)
    paths_of = by_content(files)
    ids: dict[str, str] = {}
    for fp, paths in paths_of.items():
        same_content = [r for r in rows if r["Fingerprint"] == fp]
        same_path = [r for r in rows if r["Fingerprint"] not in paths_of and set(str(r["Files"] or "").splitlines()) & set(paths)]
        row = (same_content or same_path or [None])[0]
        if row is None:
            source_id = next_id("S", [r["ID"] for r in rows])
            row = {"_row": add_row(ws, {"ID": source_id}), "ID": source_id, "Files": "", "Fingerprint": fp,
                   "Title": None, "URL": None, "Saved on": None}
            rows.append(row)
        title, url, _ = documents.get(fp, ("", "", 0))
        values = {"Files": "\n".join(paths), "Fingerprint": fp, "Status": "",
                  "Title": row["Title"] or title, "URL": row["URL"] or url, "Saved on": row["Saved on"] or saved_on.get(paths[0])}
        for header, value in values.items():
            put(ws, row["_row"], header, value)
        row["Fingerprint"] = fp
        ids[fp] = row["ID"]
    for r in rows:
        if r["Fingerprint"] not in paths_of:
            put(ws, r["_row"], "Status", "missing")
    return ids


def verify_evidence(wb: Workbook, check) -> None:
    """Set Verified on every evidence row; check(question, source, quote) returns '✓' or '✗ <reason>'."""
    ws = wb[EVIDENCE]
    for r in records(ws):
        put(ws, r["_row"], "Verified", check(r["Question"], r["Source"], r["Quote"] or ""))


# ---------------------------------------------------------------- derived columns and layout
def derive_sources(wb: Workbook, project) -> None:
    """Copy title and file from Sources into Evidence, and list each question's sources on its topic sheet.

    Removed evidence is left out of the list, so a source a reviewer removed disappears there too.
    """
    labels = project.topic_columns()
    sources = {r["ID"]: r for r in records(wb[SOURCES])}
    evidence = records(wb[EVIDENCE])
    for r in evidence:
        source = sources.get(r["Source"], {})
        put(wb[EVIDENCE], r["_row"], "Title", source.get("Title"))
        put(wb[EVIDENCE], r["_row"], "File", source.get("Files"))
    for q in questions(wb, project):
        used = dict.fromkeys(r["Source"] for r in evidence if r["Question"] == q["id"] and not r["Removed"])
        lines = [" — ".join(filter(None, (sources.get(s, {}).get("Title"), sources.get(s, {}).get("URL")))) or s for s in used]
        put(wb[q["_sheet"]], q["_row"], labels["sources"], "\n".join(lines) or None)


def derive_overview(wb: Workbook, project) -> None:
    """One Overview row per topic, with Excel formulas, so the counts stay right after edits in Excel."""
    labels = project.topic_columns()
    overview = wb[OVERVIEW]
    overview.delete_rows(2, overview.max_row)
    reads = topics(wb)
    for sheet in topic_sheets(wb):
        columns = headers(wb[sheet])
        col = {key: f"'{sheet}'!{get_column_letter(columns[labels[key]])}:{get_column_letter(columns[labels[key]])}"
               for key in ("question", "answer", "status")}
        row = overview.max_row + 1
        values = (sheet, reads.get(sheet, ""), f"=COUNTA({col['question']})-1", f"=COUNTA({col['answer']})-1",
                  *(f'=COUNTIF({col["status"]},"{s}")' for s in STATUSES), f"=D{row}-E{row}-F{row}-G{row}")
        for column, value in enumerate(values, start=1):
            overview.cell(row, column).value = value


def format_all(wb: Workbook, project) -> None:
    """Sheets in a fixed order, each laid out as a table; topic sheets also get statuses and hidden helper columns."""
    order = [OVERVIEW, *topic_sheets(wb), TOPICS, EVIDENCE, SOURCES, RUNS]
    for position, name in enumerate(order):
        wb.move_sheet(name, offset=position - wb.sheetnames.index(name))
    labels = project.topic_columns()
    derived = {labels["sources"], "Title", "File"} | set(FIXED[OVERVIEW])
    for number, ws in enumerate(wb.worksheets, start=1):
        is_topic = ws.title not in FIXED
        lay_out(ws, f"T{number}_{re.sub(r'[^A-Za-z0-9]', '_', ws.title)}",
                derived if ws.title in (OVERVIEW, EVIDENCE) or is_topic else set())
        if is_topic:
            columns = headers(ws)
            for hidden in ("asked", "scope"):
                ws.column_dimensions[get_column_letter(columns[hidden])].hidden = True
            mark_statuses(ws, get_column_letter(columns[labels["status"]]))
            ws.freeze_panes = "C2"


def lay_out(ws, table_name: str, derived: set) -> None:
    """Header style (grey for derived columns), wrapped text, fitted widths, frozen header, one Excel table."""
    for name in list(ws.tables):
        del ws.tables[name]
    last_row, last_col = max(ws.max_row, 2), max(ws.max_column, 1)
    for cell in ws[1]:
        grey = cell.value in derived
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7F7F7F" if grey else "1F3864")
        cell.comment = Comment(DERIVED_NOTE, "kb-builder") if grey and ws.title != OVERVIEW else None
    for column in range(1, last_col + 1):
        texts = (str(ws.cell(row, column).value or "") for row in range(1, last_row + 1))
        longest = max((len(line) for text in texts if not text.startswith("=") for line in text.splitlines()), default=0)
        ws.column_dimensions[get_column_letter(column)].width = min(max(longest + 2, 10), 60)
        for row in range(2, last_row + 1):
            ws.cell(row, column).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    table = Table(displayName=table_name[:250], ref=f"A1:{get_column_letter(last_col)}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)


def mark_statuses(ws, letter: str) -> None:
    """In the status column: a drop-down with the three statuses, and a colour per status."""
    cells = f"{letter}2:{letter}{max(ws.max_row, 2)}"
    ws.data_validations.dataValidation = []
    rule = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"', allow_blank=True)
    rule.add(cells)
    ws.add_data_validation(rule)
    ws.conditional_formatting = ConditionalFormattingList()
    for status, colour in zip(STATUSES, ("C6EFCE", "FFEB9C", "FFC7CE")):
        ws.conditional_formatting.add(cells, CellIsRule(operator="equal", formula=[f'"{status}"'],
                                                        fill=PatternFill("solid", bgColor=colour)))
